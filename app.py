import os
import io
import json
import time
from datetime import date
import streamlit as st
import pandas as pd
from dotenv import load_dotenv
from google import genai
from google.genai import types
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

load_dotenv()


def get_secret(name):
    try:
        return st.secrets[name]
    except Exception:
        return os.getenv(name)


client = genai.Client(api_key=get_secret("GEMINI_API_KEY"))

CATEGORIES = {
    "Building": [
        "Excavation", "PCC", "RCC (Concrete)", "Steel / Reinforcement",
        "Brickwork", "Plaster", "Flooring / Tiles", "Paint",
        "Doors & Windows", "Roofing", "Plumbing", "Electrical",
    ],
    "Road": [
        "Earthwork", "Sub-base", "Base Course", "Asphalt / Carpeting",
        "Kerb & Channel", "Drains / Culverts", "Road Marking",
    ],
}


def ask_ai(contents):
    models = ["gemini-2.5-flash", "gemini-2.5-flash-lite"]
    last_error = None
    for model in models:
        for attempt in range(3):
            try:
                return client.models.generate_content(
                    model=model,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json"
                    ),
                )
            except Exception as e:
                last_error = e
                time.sleep(3 * (attempt + 1))
    raise last_error


def make_excel(info, df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quantities"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "NAQVI QTY PRO - Bill of Quantities"
    ws["A1"].font = Font(size=16, bold=True, color="1F3864")
    ws.merge_cells("A1:H1")
    ws["A2"] = "Quantities only - Rates to be filled by user"
    ws["A2"].font = Font(size=11, italic=True)
    ws.merge_cells("A2:H2")

    row = 4
    for label, value in info:
        ws.cell(row=row, column=2, value=label).font = Font(bold=True)
        ws.cell(row=row, column=3, value=value)
        row += 1

    head_row = row + 1
    headers = ["S.No", "Item", "Description", "Unit", "Qty", "Rate", "Amount", "Remarks"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=head_row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    r = head_row + 1
    for i, rec in enumerate(df.to_dict("records"), 1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=rec.get("Item", ""))
        ws.cell(row=r, column=3, value=rec.get("Description", ""))
        ws.cell(row=r, column=4, value=rec.get("Unit", ""))
        ws.cell(row=r, column=5, value=rec.get("Qty", 0))
        ws.cell(row=r, column=6, value=None)
        ws.cell(row=r, column=7, value=f"=E{r}*F{r}")
        ws.cell(row=r, column=8, value=rec.get("Remarks", ""))
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = border
        r += 1

    ws.cell(row=r, column=6, value="TOTAL").font = Font(bold=True)
    ws.cell(row=r, column=7, value=f"=SUM(G{head_row + 1}:G{r - 1})").font = Font(bold=True)

    for col, w in zip("ABCDEFGH", [7, 24, 40, 10, 12, 12, 14, 35]):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


st.set_page_config(page_title="Naqvi QTY Pro", page_icon="🏗️", layout="wide")


def show_photo():
    if os.path.exists("friends.jpg"):
        left, mid, right = st.columns([1, 2, 1])
        with mid:
            st.image("friends.jpg", use_container_width=True)
            st.markdown(
                "<p style='text-align:center'><b>Shoaib &amp; Ali</b> - Friends forever 🤝</p>",
                unsafe_allow_html=True,
            )


if not st.session_state.get("ok"):
    show_photo()
    st.title("🏗️ Naqvi QTY Pro")
    st.caption("A gift from Shoaib, Master Five Architecture & Survey")
    pw = st.text_input("Access code", type="password")
    if pw and pw == get_secret("APP_PASSWORD"):
        st.session_state["ok"] = True
        st.rerun()
    elif pw:
        st.error("Incorrect access code.")
    st.stop()

show_photo()
st.title("🏗️ Naqvi QTY Pro")
st.caption("A gift from Shoaib, Master Five Architecture & Survey")
st.info("Always verify AI-generated quantities yourself. Results are not final.")

st.header("1. Client Details")
c1, c2, c3 = st.columns(3)
client_name = c1.text_input("Client Name *")
client_phone = c2.text_input("Phone / WhatsApp")
client_city = c3.text_input("Client City")

st.header("2. Project Details")
p1, p2, p3 = st.columns(3)
project_name = p1.text_input("Project Name *")
location = p2.text_input("Location / Address *")
project_date = p3.date_input("Date", date.today())
q1, q2 = st.columns(2)
project_type = q1.radio("Project Type", list(CATEGORIES.keys()), horizontal=True)
plot_size = q2.text_input("Plot Size / Road Length", placeholder="e.g. 10 Marla or 2.5 km")
notes = st.text_area("Additional Notes (optional)", height=70)

st.header("3. Upload Drawing")
f = st.file_uploader("PDF or Image", type=["pdf", "png", "jpg", "jpeg"])

if f:
    st.header("4. Select Quantities Required")
    items = st.multiselect(
        "Choose items",
        CATEGORIES[project_type],
        default=CATEGORIES[project_type][:3],
    )
    extra = st.text_input("Other items (optional)", placeholder="e.g. boundary wall, septic tank")

    if st.button("Extract Quantities", type="primary"):
        if not client_name or not project_name or not location:
            st.warning("Please fill in Client Name, Project Name and Location.")
        elif not items and not extra:
            st.warning("Please select at least one item.")
        else:
            wanted = ", ".join(items + ([extra] if extra else []))
            prompt = f"""You are an expert quantity surveyor.
Project: {project_name}, Location: {location}, Type: {project_type}, Size: {plot_size}.
Notes: {notes}
From this drawing/text, extract quantities ONLY for: {wanted}.
Do not include any other items. No rates.
Return a JSON list:
[{{"Item":"","Description":"","Unit":"","Qty":0,"Remarks":""}}]
If a dimension is missing, put Qty 0 and explain in Remarks."""
            with st.spinner("AI is reading the drawing... (may take up to 20 seconds)"):
                try:
                    r = ask_ai([
                        types.Part.from_bytes(data=f.getvalue(), mime_type=f.type),
                        prompt,
                    ])
                    st.session_state["df"] = pd.DataFrame(json.loads(r.text))
                    st.session_state["info"] = [
                        ("Client", client_name),
                        ("Phone", client_phone),
                        ("City", client_city),
                        ("Project", project_name),
                        ("Location", location),
                        ("Type", project_type),
                        ("Size", plot_size),
                        ("Date", project_date.strftime("%d-%m-%Y")),
                    ]
                except Exception as e:
                    st.error(f"Error: {e}")

if "df" in st.session_state:
    st.header("5. Quantities (review and edit)")
    edited = st.data_editor(st.session_state["df"], num_rows="dynamic", use_container_width=True)
    st.download_button(
        "📥 Download Excel",
        make_excel(st.session_state["info"], edited),
        f"BOQ_{project_name or 'project'}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )